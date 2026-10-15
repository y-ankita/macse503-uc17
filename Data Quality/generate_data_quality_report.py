
from pathlib import Path
from datetime import datetime
import getpass

import pandas as pd
import psycopg2
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# =====================================================
# YOUR SPRINT 2 PROJECT LOCATION
# =====================================================

BASE_DIR = Path(
    r"C:\Users\maaha\OneDrive\Desktop\VIT\DATA ENGINEERING\DA\SPRINT 2"
)

OUTPUT_DIR = BASE_DIR / "data_quality_output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

DATASETS = {
    "access": ("cleaned", "access_clean",
               "standardized", "access_standardized"),
    "attendance": ("cleaned", "attendance_clean",
                  "standardized", "attendance_standardized"),
    "hrms": ("cleaned", "hrms_clean",
             "standardized", "hrms_standardized"),
    "leave": ("cleaned", "leave_clean",
              "standardized", "leave_standardized"),
    "performance": ("cleaned", "performance_clean",
                   "standardized", "performance_standardized"),
}

EVIDENCE_FILES = [
    "EHCMDP_Data_Quality_Analysis.xlsx",
    "Staging_Diagnostic_Report.xlsx",
    "Cleaned_Data_Validation.xlsx",
    "Standardization_Evidence.xlsx",
    "Standardized_Data_Validation.xlsx",
]


# =====================================================
# DATABASE CONNECTION
# =====================================================

def connect_database():
    password = getpass.getpass("Enter PostgreSQL password: ")

    return psycopg2.connect(
        host="localhost",
        port=5432,
        database="EHCMDP_DB",
        user="postgres",
        password=password,
    )


# =====================================================
# DATABASE QUALITY METRICS
# =====================================================

def get_table_metrics(conn, schema, table):
    with conn.cursor() as cur:
        cur.execute("""
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = %s
              AND table_name = %s
            ORDER BY ordinal_position
        """, (schema, table))

        columns = [row[0] for row in cur.fetchall()]

        if not columns:
            raise ValueError(
                f"Table not found: {schema}.{table}"
            )

        qualified_table = f'"{schema}"."{table}"'

        cur.execute(f"SELECT COUNT(*) FROM {qualified_table}")
        row_count = cur.fetchone()[0]

        # Count SQL NULL cells.
        missing_sql = ", ".join(
            f'COUNT(*) FILTER (WHERE "{col}" IS NULL)'
            for col in columns
        )
        cur.execute(f"SELECT {missing_sql} FROM {qualified_table}")
        missing_cells = sum(cur.fetchone())

        # Count exact duplicate rows beyond the first occurrence.
        column_list = ", ".join(f'"{col}"' for col in columns)
        cur.execute(f"""
            SELECT COALESCE(SUM(n - 1), 0)
            FROM (
                SELECT COUNT(*) AS n
                FROM {qualified_table}
                GROUP BY {column_list}
                HAVING COUNT(*) > 1
            ) AS duplicate_groups
        """)
        duplicate_rows = int(cur.fetchone()[0])

        # Count rows where every column is NULL.
        empty_condition = " AND ".join(
            f'"{col}" IS NULL' for col in columns
        )
        cur.execute(f"""
            SELECT COUNT(*)
            FROM {qualified_table}
            WHERE {empty_condition}
        """)
        empty_rows = cur.fetchone()[0]

    return {
        "rows": row_count,
        "columns": len(columns),
        "missing_cells": missing_cells,
        "duplicate_rows": duplicate_rows,
        "fully_empty_rows": empty_rows,
    }


def collect_results(conn):
    records = []

    for dataset, names in DATASETS.items():
        cleaned_schema, cleaned_table, std_schema, std_table = names

        cleaned = get_table_metrics(
            conn, cleaned_schema, cleaned_table
        )
        standardized = get_table_metrics(
            conn, std_schema, std_table
        )

        records.append({
            "dataset": dataset,
            "cleaned_rows": cleaned["rows"],
            "standardized_rows": standardized["rows"],
            "cleaned_columns": cleaned["columns"],
            "standardized_columns": standardized["columns"],
            "row_counts_match":
                cleaned["rows"] == standardized["rows"],
            "column_counts_match":
                cleaned["columns"] == standardized["columns"],
            "standardized_missing_cells":
                standardized["missing_cells"],
            "standardized_duplicate_rows":
                standardized["duplicate_rows"],
            "standardized_fully_empty_rows":
                standardized["fully_empty_rows"],
        })

    return pd.DataFrame(records)


# =====================================================
# EXCEL REPORT
# =====================================================

def create_excel_report(results):
    output = OUTPUT_DIR / "Final_Data_Quality_Summary.xlsx"

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        results.to_excel(
            writer, sheet_name="Validation Summary", index=False
        )

        evidence = pd.DataFrame({
            "Evidence file": EVIDENCE_FILES,
            "Exists in output folder": [
                (OUTPUT_DIR / filename).exists()
                for filename in EVIDENCE_FILES
            ],
        })
        evidence.to_excel(
            writer, sheet_name="Evidence Checklist", index=False
        )

        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

            for cells in sheet.columns:
                width = max(
                    len(str(cell.value or "")) for cell in cells
                )
                sheet.column_dimensions[
                    cells[0].column_letter
                ].width = min(width + 2, 42)

    return output


# =====================================================
# WORD REPORT HELPERS
# =====================================================

def shade_cell(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(heading)
        shade_cell(cell, "1F4E78")

        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    return table


# =====================================================
# GENERATE DATA QUALITY REPORT
# =====================================================

def create_word_report(results):
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9)

    title = doc.add_heading("Sprint 2: Data Quality Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(
        "Data Quality Analysis | Cleaning and Standardization | Validation"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"Generated on: {datetime.now().strftime('%d %B %Y %H:%M')}"
    )

    doc.add_heading("1. Executive Summary", level=1)

    total_rows = int(results["standardized_rows"].sum())
    total_missing = int(
        results["standardized_missing_cells"].sum()
    )

    doc.add_paragraph(
        "This report covers five workforce datasets: access, attendance, "
        "hrms, leave, and performance. It summarizes data cleaning, "
        "standardization, and validation performed on their PostgreSQL "
        "tables."
    )
    doc.add_paragraph(
        f"The standardized tables contain {total_rows:,} rows and "
        f"{total_missing:,} SQL NULL cells. Dataset-level validation "
        "results are presented below."
    )

    doc.add_heading("2. Scope and Database Tables", level=1)

    add_table(
        doc,
        ["Dataset", "Cleaned table", "Standardized table"],
        [
            [
                name,
                f"cleaned.{tables[1]}",
                f"standardized.{tables[3]}",
            ]
            for name, tables in DATASETS.items()
        ],
    )

    doc.add_heading("3. Cleaning and Standardization", level=1)
    doc.add_paragraph(
        "The cleaning process created copies of the staging data and "
        "removed fully empty rows and exact duplicate rows according "
        "to the cleaning script. Staging source tables were not modified."
    )
    doc.add_paragraph(
        "The standardization process created standardized copies, "
        "normalized text whitespace, and converted blank text strings "
        "to SQL NULL where applicable. A successful run does not prove "
        "that all business rules or field formats are consistent."
    )

    doc.add_heading("4. Validation Results", level=1)

    validation_rows = []
    for _, row in results.iterrows():
        validation_rows.append([
            row["dataset"],
            f'{int(row["standardized_rows"]):,}',
            int(row["standardized_columns"]),
            "Yes" if row["row_counts_match"] else "No",
            "Yes" if row["column_counts_match"] else "No",
            int(row["standardized_duplicate_rows"]),
            int(row["standardized_fully_empty_rows"]),
            f'{int(row["standardized_missing_cells"]):,}',
        ])

    add_table(
        doc,
        [
            "Dataset", "Rows", "Columns", "Rows match",
            "Columns match", "Duplicate rows", "Empty rows",
            "Missing cells",
        ],
        validation_rows,
    )

    doc.add_paragraph(
        "Matching row and column counts indicate that standardization "
        "preserved the cleaned table shape. Duplicate and empty-row "
        "checks cover exact duplicates and rows containing SQL NULL "
        "in every column; they do not establish complete semantic validity."
    )

    doc.add_heading("5. Missing-Data Findings", level=1)

    hrms = results[results["dataset"] == "hrms"].iloc[0]

    doc.add_paragraph(
        f"The standardized HRMS table has "
        f'{int(hrms["standardized_missing_cells"]):,} SQL NULL cells. '
        "Existing diagnostics identified high missingness in fields "
        "including manager_id, last_promotion_date, termination_date, "
        "and last_training_date. Investigate the source and ingestion "
        "process before deciding whether values should be filled."
    )
    doc.add_paragraph(
        "Missing-cell counts refer to SQL NULL values. Empty strings "
        "or placeholder text may require separate checks."
    )

    doc.add_heading("6. Limitations", level=1)
    for item in [
        "Exact duplicate checks do not identify near-duplicate employee records.",
        "SQL NULL counts do not automatically identify empty strings or placeholder values.",
        "Matching row counts do not prove every field value is correct.",
        "Whitespace normalization does not validate all dates, categories, IDs, or numeric ranges.",
        "HRMS missingness remains unresolved; its cause has not been confirmed.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("7. Recommendations", level=1)
    for item in [
        "Investigate HRMS source completeness and ingestion patterns.",
        "Confirm whether highly incomplete HRMS fields are expected.",
        "Define field-level checks for dates, categories, identifiers, and numeric ranges.",
        "Do not invent missing values or fill business-critical fields without an approved rule.",
        "Retain the existing analysis, cleaning, and validation evidence.",
    ]:
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("8. Evidence Checklist", level=1)
    add_table(
        doc,
        ["Evidence file", "Status in data_quality_output"],
        [
            [
                filename,
                "Found" if (OUTPUT_DIR / filename).exists()
                else "Not found",
            ]
            for filename in EVIDENCE_FILES
        ],
    )

    doc.add_paragraph(
        "Assigned deliverables: Data Quality Report, "
        "cleaned/standardized data, and validation evidence."
    )

    doc.add_heading("9. Conclusion", level=1)
    doc.add_paragraph(
        "The five datasets have been processed through the assigned "
        "cleaning, standardization, and validation workflow. The "
        "validation results describe the checks performed, while the "
        "missing-data issue in HRMS remains a priority for investigation. "
        "The report does not claim that all possible data quality issues "
        "have been eliminated."
    )

    output = OUTPUT_DIR / "Final_Data_Quality_Report.docx"
    doc.save(output)
    return output


# =====================================================
# MAIN
# =====================================================

def main():
    conn = None

    try:
        print("Connecting to EHCMDP_DB...")
        conn = connect_database()

        print("Validating the five datasets...")
        results = collect_results(conn)

        word_path = create_word_report(results)
        excel_path = create_excel_report(results)

        print("\nSUCCESS: Reports generated.")
        print(f"Word report:  {word_path}")
        print(f"Excel report: {excel_path}")

        print("\nValidation summary:")
        print(results.to_string(index=False))

    except Exception as error:
        print("\nERROR:", error)
        print(
            "Check your PostgreSQL connection, table names, "
            "permissions, and installed libraries."
        )
        raise

    finally:
        if conn is not None:
            conn.close()


if __name__ == "__main__":
    main()